#!/usr/bin/env python3
"""Convert both EP1 Excel databases → drama-quotes-data.js with SRT-synced YouTube times."""

import json
import re
from pathlib import Path

from openpyxl import load_workbook

BASE = Path(__file__).resolve().parent
OUT_JS = BASE.parent / "drama-quotes-data.js"
YOUTUBE_VIDEO_ID = "QJhEkjmgwis"

DATASETS = [
    {
        "xlsx": BASE / "台語資料庫_EP1.xlsx",
        "srt": BASE / "ep1前半段字幕台語字幕.srt",
        "mandarin_srt": BASE / "ep1前半段字幕拜託.srt",
        "label": "EP1前半",
    },
    {
        "xlsx": BASE / "台語資料庫_EP1後半.xlsx",
        "srt": BASE / "_MConverter.eu_《夜市人生》第一集後半段完整190句台語漢字字幕檔 (1).srt",
        "mandarin_srt": None,
        "label": "EP1後半",
    },
]

SYNTAX_IDS = {
    "kā-disposal", "kā-target", "hōo-passive", "hōo-causative",
    "comparative", "potential", "extent-comp", "conditionals",
}
SEMANTICS_IDS = {
    "aspect-perf", "aspect-cont", "aspect-neg", "modal-neg",
    "diminutive-a", "directional",
}
PHONOLOGY_IDS = {
    "neutral-tone", "contraction", "reduplication", "checked-tone", "literary-coll",
}
PRAGMATICS_IDS = {
    "prag-angry", "prag-sad", "prag-happy", "prag-sarcasm", "prag-shock", "prag-idiom",
}

PRAG_META = {
    "prag-angry": ("生氣", "吵架", []),
    "prag-sad": ("悲傷", "討債", []),
    "prag-happy": ("快樂", "夜市生意", []),
    "prag-sarcasm": ("冷淡", "親密揶揄", []),
    "prag-shock": ("疑惑", "理性討論", []),
    "prag-idiom": ("平靜", "日常對話交流", ["諺語"]),
}

IMAGE_BG = (
    "https://images.unsplash.com/photo-1513885045260-6b358ae178c7"
    "?auto=format&fit=crop&w=900&q=80"
)

MAX_CUE_DURATION = 25.0


def normalize_text(s: str) -> str:
    return re.sub(r"[\s…\.。，,、！!？?「」『』""\"'\-]+", "", s or "")


def fix_time_line(line: str) -> str:
    line = line.strip()
    line = re.sub(r"(\d{2}:\d{2})\s+(\d{2},)", r"\1:\2", line)
    line = re.sub(r"\s+", "", line)
    return line.replace(".", ",")


def parse_time_range(range_str: str) -> tuple[int, int, str]:
    m = re.match(
        r"(\d{2}):(\d{2}):(\d{2})-(\d{2}):(\d{2}):(\d{2})",
        range_str or "",
    )
    if not m:
        return 0, 0, ""
    sh, sm, ss, eh, em, es = map(int, m.groups())
    start = sh * 3600 + sm * 60 + ss
    end = eh * 3600 + em * 60 + es
    ts_range = f"{sh:02d}:{sm:02d}:{ss:02d}-{eh:02d}:{em:02d}:{es:02d}"
    return start, end, ts_range


def parse_srt(path: Path) -> list[dict]:
    raw = path.read_text(encoding="utf-8")
    raw = re.sub(r"(\d{2}:\d{2})\s+(\d{2},)", r"\1:\2", raw)
    lines = raw.splitlines()
    entries = []
    i = 0
    while i < len(lines):
        while i < len(lines) and not lines[i].strip():
            i += 1
        if i >= len(lines):
            break
        if re.match(r"^\d+$", lines[i].strip()):
            i += 1
        if i >= len(lines):
            break
        time_line = fix_time_line(lines[i])
        i += 1
        if "-->" not in time_line:
            continue
        text_lines = []
        while i < len(lines):
            s = lines[i].strip()
            if not s:
                i += 1
                break
            nxt = fix_time_line(lines[i + 1]) if i + 1 < len(lines) else ""
            if re.match(r"^\d+$", s) and "-->" in nxt:
                break
            text_lines.append(s)
            i += 1
        m = re.search(
            r"(\d{2}):(\d{2}):(\d{2}),(\d{3})-->(\d{2}):(\d{2}):(\d{2}),(\d{3})",
            time_line,
        )
        if not m or not text_lines:
            continue
        sh, sm, ss, sms, eh, em, es, ems = map(int, m.groups())
        start = sh * 3600 + sm * 60 + ss + sms / 1000
        end = eh * 3600 + em * 60 + es + ems / 1000
        if end < start:
            end = start + 1.5
        if end - start > MAX_CUE_DURATION:
            start = max(end - 3.0, 0)
        text = " ".join(text_lines).strip()
        entries.append({
            "text": text,
            "norm": normalize_text(text),
            "start": start,
            "end": end,
            "timestamp": time_line,
            "timestamp_range": (
                f"{sh:02d}:{sm:02d}:{ss:02d}-{eh:02d}:{em:02d}:{es:02d}"
            ),
        })
    return entries


def build_srt_lookup(entries: list[dict]) -> dict[str, dict]:
    lookup = {}
    for e in entries:
        lookup[e["norm"]] = e
        lookup[e["text"]] = e
    return lookup


def align_mandarin(taigi_start: float, mandarin_entries: list[dict]) -> str:
    for m in mandarin_entries:
        if abs(m["start"] - taigi_start) < 0.6:
            return m["text"]
    return ""


def seconds_to_episode_label(sec: float) -> str:
    s = int(sec)
    h, rem = divmod(s, 3600)
    m, ss = divmod(rem, 60)
    if h:
        return f"第 1 集 {h:02d}:{m:02d}:{ss:02d}"
    return f"第 1 集 {m:02d}:{ss:02d}"


def split_ids(raw: str) -> list[str]:
    if not raw:
        return []
    return [x.strip() for x in str(raw).replace(",", "、").split("、") if x.strip()]


def tags_for_category(tag_ids: list[str], id_set: set[str]) -> str:
    return ", ".join(t for t in tag_ids if t in id_set) or None


def infer_meta(tag_ids: list[str]) -> tuple[str, str, list[str]]:
    for tid in tag_ids:
        if tid in PRAGMATICS_IDS and tid in PRAG_META:
            return PRAG_META[tid]
    return "平靜", "日常對話交流", ["一般交流用句"]


def make_practice(quote: str) -> str:
    snippet = quote if len(quote) <= 18 else quote[:18] + "…"
    return f"請朗讀並標註台羅：「{snippet}」"


def js_str(s: str) -> str:
    return json.dumps(s, ensure_ascii=False)


def resolve_timing(quote: str, time_range: str, ts_raw: str, srt_lookup: dict) -> dict:
    srt = srt_lookup.get(normalize_text(quote)) or srt_lookup.get(quote.strip())
    if srt:
        yt_start = int(srt["start"])
        yt_end = max(int(srt["end"]), yt_start + 1)
        return {
            "youtubeStart": yt_start,
            "youtubeEnd": yt_end,
            "timestampRange": srt["timestamp_range"],
            "timestampRaw": srt["timestamp"],
            "timeSource": "srt",
        }

    ex_start, ex_end, ex_range = parse_time_range(time_range)
    if ex_start or ex_end:
        return {
            "youtubeStart": ex_start,
            "youtubeEnd": ex_end or ex_start + 2,
            "timestampRange": ex_range or time_range,
            "timestampRaw": ts_raw or "",
            "timeSource": "excel",
        }

    return None


def load_dataset(cfg: dict) -> list[dict]:
    if not cfg["xlsx"].exists():
        raise SystemExit(f"Missing {cfg['xlsx']}")

    srt_entries = parse_srt(cfg["srt"])
    srt_lookup = build_srt_lookup(srt_entries)
    mandarin_srt = parse_srt(cfg["mandarin_srt"]) if cfg["mandarin_srt"] else []

    wb = load_workbook(cfg["xlsx"])
    ws = wb.active
    rows = [r for r in ws.iter_rows(min_row=2, values_only=True) if r and r[3]]

    entries = []
    missing_time = []
    for r in rows:
        (
            _seq, ts_raw, time_range, quote, pinyin, translation,
            _syntax_col, _semantics_col, _phonology_col, _prag_col,
            grammar_labels, tag_ids_raw, explanation,
        ) = r[:13]

        timing = resolve_timing(quote, time_range, ts_raw or "", srt_lookup)
        if not timing:
            missing_time.append(quote)
            continue

        tag_ids = split_ids(tag_ids_raw)
        tone, context, sentence_types = infer_meta(tag_ids)
        yt_start = timing["youtubeStart"]

        mandarin = (translation or "").strip()
        if not mandarin and mandarin_srt:
            mandarin = align_mandarin(yt_start, mandarin_srt)
        if not mandarin:
            mandarin = "（待補華語）"

        exp = (explanation or "").strip()
        grammar = grammar_labels or ""
        if not exp and grammar:
            exp = f"文法考點：{grammar}"

        entries.append({
            "drama": "夜市人生",
            "episode": 1,
            "character": "劇中角色",
            "speakerCode": cfg["label"],
            "language": "台語",
            "quote": quote,
            "pinyin": pinyin or "",
            "translation": mandarin,
            "difficulty": "Intermediate",
            "sentenceTypes": sentence_types,
            "tone": tone,
            "context": context,
            "grammar": grammar,
            "linguisticTags": tag_ids,
            "syntaxTag": tags_for_category(tag_ids, SYNTAX_IDS),
            "semanticsTags": tags_for_category(tag_ids, SEMANTICS_IDS),
            "phonologyTags": tags_for_category(tag_ids, PHONOLOGY_IDS),
            "pragmaticsTag": tags_for_category(tag_ids, PRAGMATICS_IDS),
            "explanation": exp,
            "deepDive": exp,
            "funFact": "（尚無 Fun Fact 資料）",
            "practice": make_practice(quote),
            "activity": f"兩人一組：一人以「{tone}」語氣唸出台詞，另一人用華語解釋意思。",
            "imageBg": IMAGE_BG,
            "videoUrl": f"https://www.youtube.com/watch?v={YOUTUBE_VIDEO_ID}&t={yt_start}",
            "youtubeStart": yt_start,
            "youtubeEnd": timing["youtubeEnd"],
            "audioStart": yt_start,
            "audioEnd": timing["youtubeEnd"],
            "episodeTime": seconds_to_episode_label(yt_start),
            "timestampRange": timing["timestampRange"],
            "timestampRaw": timing["timestampRaw"],
            "warning": None,
            "_timeSource": timing["timeSource"],
            "_segment": cfg["label"],
        })

    if missing_time:
        print(f"Warning [{cfg['label']}]: no timing for:", missing_time)

    return entries


def main():
    all_entries = []
    for cfg in DATASETS:
        batch = load_dataset(cfg)
        print(f"{cfg['label']}: {len(batch)} rows")
        all_entries.extend(batch)

    all_entries.sort(key=lambda e: (e["youtubeStart"], e["_segment"], e["quote"]))
    for i, e in enumerate(all_entries, 1):
        e["id"] = i
        e.pop("_timeSource", None)
        e.pop("_segment", None)

    lines = [
        "// Auto-generated from 台語資料庫_EP1.xlsx + 台語資料庫_EP1後半.xlsx",
        "// YouTube timestamps synced to subtitle SRT (absolute seconds in EP001)",
        f"// Video: https://www.youtube.com/watch?v={YOUTUBE_VIDEO_ID}",
        "const dramaQuotesDatabase = [",
    ]
    for i, e in enumerate(all_entries):
        lines.append("    {")
        for key, val in e.items():
            if val is None:
                lines.append(f"        {key}: null,")
            elif isinstance(val, bool):
                lines.append(f"        {key}: {str(val).lower()},")
            elif isinstance(val, (int, float)):
                lines.append(f"        {key}: {val},")
            elif isinstance(val, list):
                items = ", ".join(js_str(x) for x in val)
                lines.append(f"        {key}: [{items}],")
            else:
                lines.append(f"        {key}: {js_str(val)},")
        lines.append("    }" + ("," if i < len(all_entries) - 1 else ""))
    lines.append("];")
    lines.append("")

    OUT_JS.write_text("\n".join(lines), encoding="utf-8")
    print(f"Wrote {len(all_entries)} entries to {OUT_JS}")
    print(f"  First: t={all_entries[0]['youtubeStart']}s — {all_entries[0]['quote'][:24]}")
    print(f"  Last:  t={all_entries[-1]['youtubeStart']}s — {all_entries[-1]['quote'][:24]}")


if __name__ == "__main__":
    main()
